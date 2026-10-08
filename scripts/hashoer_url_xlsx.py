# -*- coding: utf-8 -*-
"""hashoer 多语言子目录 URL 清单 Excel（同 druomu 格式）"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

LANGS = [
    ('en', '英文（默认）', '/'),
    ('zh', '中文', '/zh/'),
    ('ja', '日语', '/ja/'),
    ('ko', '韩语', '/ko/'),
    ('ru', '俄语', '/ru/'),
    ('fr', '法语', '/fr/'),
    ('es', '西班牙语', '/es/'),
    ('pt', '葡萄牙语', '/pt/'),
    ('th', '泰语', '/th/'),
    ('vi', '越南语', '/vi/'),
    ('bn', '孟加拉语', '/bn/'),
    ('ur', '乌尔都语', '/ur/'),
    ('fa', '波斯语', '/fa/'),
    ('ar', '阿拉伯语', '/ar/'),
    ('de', '德语', '/de/'),
    ('it', '意大利语', '/it/'),
]
BASE = 'https://hashoer.com'

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '多语言URL清单'

headers = ['#', '语言', '语言代码', '首页 URL', '产品页 URL', '隐私页 URL']
ws.append(headers)

for i, (code, name, prefix) in enumerate(LANGS, 1):
    if code == 'en':
        home, prod, priv = BASE + '/', BASE + '/products', BASE + '/privacy'
    else:
        home, prod, priv = BASE + prefix, BASE + prefix + 'products', BASE + prefix + 'privacy'
    ws.append([i, name, code, home, prod, priv])

# 样式：蓝底白字表头、英文行加粗、冻结首行、URL 蓝色
header_fill = PatternFill('solid', fgColor='1F4E79')
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF', size=11)
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center')

for row in ws.iter_rows(min_row=2):
    is_en = row[2].value == 'en'
    for c in row:
        c.alignment = Alignment(vertical='center')
        if is_en:
            c.font = Font(bold=True)
        if c.column >= 4:
            c.font = Font(color='0563C1', underline='single', bold=is_en)

ws.column_dimensions['A'].width = 5
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 10
for col in ('D', 'E', 'F'):
    ws.column_dimensions[col].width = 38
ws.freeze_panes = 'A2'

# 第二个 sheet：SEO 提交说明
ws2 = wb.create_sheet('SEO提交说明')
notes = [
    ['项目', '内容'],
    ['一键提交', '把 https://hashoer.com/sitemap.xml 提交到 Google Search Console，即覆盖全部 48 个 URL（含 hreflang 关联）'],
    ['hreflang', '每个页面已配置 16 语言互链 + x-default（默认英文），Google 会按访客地区/语言展示对应版本'],
    ['默认语言', '打开 hashoer.com（不带子目录）永远显示英文；不开 IP 自动跳转'],
    ['URL 规则', '英文在根目录；其他语言在 /语言代码/ 子目录（如 /es/ = 西班牙语）'],
    ['生效状态', '全部 URL 已上线（2026-09-29），抽查均返回 200'],
]
for r in notes:
    ws2.append(r)
for c in ws2[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = header_fill
ws2.column_dimensions['A'].width = 14
ws2.column_dimensions['B'].width = 95
for row in ws2.iter_rows(min_row=2):
    row[1].alignment = Alignment(wrap_text=True, vertical='top')

out = r'C:\Users\Administrator\WorkBuddy\建站搭子\hashoer-en\hashoer_多语言子目录URL清单.xlsx'
wb.save(out)
print('saved:', out)
